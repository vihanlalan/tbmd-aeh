# -*- coding: utf-8 -*-
"""
build_paper.py
--------------
Regenerates the manuscript (.docx) from outputs/paper/*.csv so every number
in the text and tables comes straight from run_paper_analysis.py.

    python src/run_paper_analysis.py
    python src/build_paper.py
"""

import copy
import json
import os
import pandas as pd
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

ROOT = os.path.join(os.path.dirname(__file__), '..')
OUT = os.path.join(ROOT, 'outputs', 'paper')
DOCX = os.path.join(ROOT, 'Revised_Paper_TBMD_Operationalizing_AMH.docx')

MINUS = '−'


def csv(name, **kw):
    return pd.read_csv(os.path.join(OUT, name), **kw)


def f(x, nd=2, sign=False, pct=False):
    s = f'{x:+.{nd}f}' if sign else f'{x:.{nd}f}'
    s = s.replace('-', MINUS)
    return s + ('%' if pct else '')


def p_fmt(p):
    return '< 0.001' if p < 0.001 else f'{p:.3f}'


# ------------------------------------------------------------------
# Load results
# ------------------------------------------------------------------
synth = csv('synthetic_res_detection_summary.csv', index_col=0)
us_det = csv('us_res_detection.csv', index_col=0)
de_det = csv('de_res_detection.csv', index_col=0)
comp = csv('us_bpc_components.csv', index_col=0)
grid = csv('us_dev_grid.csv')
perf = {p: csv(f'us_performance_{p}.csv', index_col=0) for p in ('development', 'lockbox', 'full')}
diff = csv('us_rcf_minus_unconditional.csv', index_col=0)
rb = csv('us_regime_breakdown_full.csv')
cost = csv('us_cost_sensitivity.csv')
de_perf = csv('de_performance_full.csv', index_col=0)
de_rb = csv('de_regime_breakdown_full.csv')
summary = json.load(open(os.path.join(OUT, 'summary.json')))
sel = summary['selected_config']
de_diff = summary['de_rcf_minus_unconditional']
n_us = summary['universe']['us_n_stocks']
n_de = summary['universe']['de_n_stocks']

D, L, F = perf['development'], perf['lockbox'], perf['full']
rcfL, uncL, ewL, spyL = (L.loc[k] for k in ['RCF', 'Unconditional', 'Buy & Hold (equal weight)', 'Buy & Hold (SPY)'])
rcfD, uncD = D.loc['RCF'], D.loc['Unconditional']
rcfF, uncF, ewF = F.loc['RCF'], F.loc['Unconditional'], F.loc['Buy & Hold (equal weight)']
dL, dD, dF = diff.loc['lockbox'], diff.loc['development'], diff.loc['full']
vix_det = us_det.loc['VIX >= 20']
auc_range = (us_det['auc'].min(), us_det['auc'].max())
dev_best = grid['dev_rcf_sharpe'].max()
n_pos = int((grid['dev_rcf_sharpe'] > 0).sum())
cost0 = cost[(cost.cost_bps == 0) & (cost.strategy == 'RCF')].iloc[0]
de_rcf, de_unc, de_ew = de_perf.loc['RCF'], de_perf.loc['Unconditional'], de_perf.loc['Buy & Hold (equal weight)']
de_rv = de_det.loc['DAX 21d realized vol >= 20%']
de_dd = de_det.loc['DAX drawdown >= 10%']
stress = rb.iloc[2]
de_stress = de_rb.iloc[2]
zs = comp.loc['Z_sent']
zi = comp.loc['Z_illiq']
zi_lo, zi_hi = zi.iloc[:4].min(), zi.iloc[:4].max()
zh = comp.loc['Z_herd']
zb = comp.loc['BPC']
zb_lo, zb_hi = zb.iloc[:4].min(), zb.iloc[:4].max()

# ------------------------------------------------------------------
# Document scaffolding (reuse the existing document's styles)
# ------------------------------------------------------------------
doc = docx.Document(DOCX)
body = doc.element.body
list_numpr = None
for p in doc.paragraphs:
    if p.style.name == 'List Paragraph' and p._p.pPr is not None and p._p.pPr.numPr is not None:
        list_numpr = copy.deepcopy(p._p.pPr.numPr)
        break
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)


def para(text='', style='Normal', align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=None):
    p = doc.add_paragraph(style=style)
    if style == 'Normal':
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    return p


def rich(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """parts: list of (text, bold, italic)"""
    p = doc.add_paragraph(style='Normal')
    p.alignment = align
    for text, b, i in parts:
        r = p.add_run(text)
        r.bold, r.italic = b, i
    return p


def h1(t):
    doc.add_paragraph(t, style='Heading 1')


def h2(t):
    doc.add_paragraph(t, style='Heading 2')


def bullet(label, text):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if list_numpr is not None:
        pPr = p._p.get_or_add_pPr()
        pStyle = pPr.find(qn('w:pStyle'))
        if pStyle is not None:
            pStyle.addnext(copy.deepcopy(list_numpr))
        else:
            pPr.insert(0, copy.deepcopy(list_numpr))
    if label:
        p.add_run(label + ' ').bold = True
    p.add_run(text)


def equation(t):
    p = para(t, align=WD_ALIGN_PARAGRAPH.CENTER)
    for r in p.runs:
        r.font.name = 'Cambria Math'
        r.font.size = Pt(11)


def _borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'BFBFBF')
        b.append(e)
    anchor = next((tblPr.find(qn(f'w:{t}')) for t in ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption')
                   if tblPr.find(qn(f'w:{t}')) is not None), None)
    if anchor is not None:
        anchor.addprevious(b)
    else:
        tblPr.append(b)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def table(caption, header, rows, note=None, widths=None):
    cp = para(caption, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=10)
    cp.paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    _borders(t)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ''
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size = True, Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _shade(c, '2E75B6')
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            c.text = ''
            r = c.paragraphs[0].add_run(str(v))
            r.font.size = Pt(8.5)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if i % 2 == 0:
                _shade(c, 'F2F6FA')
    trPr = t.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader')
    trPr.append(hdr)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths):
                row.cells[j].width = docx.shared.Inches(w)
    if note:
        para(note, align=WD_ALIGN_PARAGRAPH.JUSTIFY, italic=True, size=8.5)
    else:
        para('')


# ------------------------------------------------------------------
# Title block
# ------------------------------------------------------------------
para('Can the Adaptive Markets Hypothesis Be Operationalized? A Pre-Specified, Out-of-Sample Test of '
     'Behavioral Regime Detection and Regime-Conditional Trading',
     align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16)
para('Vihan Lalan', align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
para('September 2026', align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=10)

h1('Abstract')
para(
    "Lo's (2004, 2017) Adaptive Markets Hypothesis (AMH) holds that market efficiency varies over time, but it "
    "does not say how to detect an inefficient regime in real time or whether doing so is economically useful. "
    "We build a fully specified operationalization and subject it to a pre-specified, out-of-sample test. A "
    "Real-Time Efficiency Score (RES) combines rolling Hurst, Lo-MacKinlay variance-ratio and Ljung-Box "
    "statistics; a Behavioral Proxy Composite (BPC) aggregates five observable proxies for herding, "
    "autocorrelation, loss aversion, sentiment and illiquidity; and a Regime-Conditional Filter (RCF) deploys a "
    f"cross-sectional long/short strategy only when the RES signals inefficiency. Using {n_us} S&P 100 "
    "constituents from 2005 to 2025, we select one of twelve pre-declared configurations on 2005–2020 data "
    "and evaluate it once on a held-out 2021–2025 lockbox, then apply it unchanged to "
    f"{n_de} DAX constituents. The RES does not detect stress regimes: its AUC against four independent regime "
    f"definitions (VIX, realized volatility, drawdowns, NBER recessions) ranges from {f(auc_range[0], 2)} to "
    f"{f(auc_range[1], 2)}, with every bootstrap confidence interval covering 0.5, and it is only weakly "
    f"informative even in a Hamilton (1989) simulation built to contain the regimes it targets (mean AUC "
    f"{f(synth.loc['mean', 'auc'], 3)} across 20 seeds). The best development-period configuration earns a "
    f"Sharpe ratio of {f(rcfD['sharpe'], 2)} (deflated Sharpe ratio {f(rcfD['dsr_n12'], 2)}); on the lockbox it "
    f"earns {f(rcfL['sharpe'], 2)} (95% CI {f(rcfL['sharpe_ci_low'], 2)} to {f(rcfL['sharpe_ci_high'], 2)}), "
    "and the filter does not reliably improve on deploying the same signal unconditionally in either market. "
    "Of the behavioral proxies, only illiquidity carries substantial, non-circular information about market "
    f"stress (AUC {f(zi_lo, 2)}–{f(zi_hi, 2)}). We conclude that this operationalization of "
    "the AMH does not survive out-of-sample testing, and we document which of its components fail and why.")

rich([('Keywords: ', True, False),
      ('adaptive markets hypothesis; market efficiency; regime detection; behavioral finance; out-of-sample '
       'testing; deflated Sharpe ratio', False, False)], align=WD_ALIGN_PARAGRAPH.LEFT)
rich([('JEL classification: ', True, False), ('G12, G14, G41, C58', False, False)], align=WD_ALIGN_PARAGRAPH.LEFT)

# ------------------------------------------------------------------
# I. Introduction
# ------------------------------------------------------------------
h1('I. Introduction')
para(
    "That markets are not uniformly efficient is no longer seriously contested. Momentum (Jegadeesh & Titman, "
    "1993), post-earnings-announcement drift (Ball & Brown, 1968), long-horizon reversal (DeBondt & Thaler, "
    "1985) and the illiquidity premium (Amihud, 2002) all indicate that prices do not instantly impound "
    "available information. At the same time, many documented anomalies fail to replicate (Hou, Xue & Zhang, "
    "2020) or decay after publication (McLean & Pontiff, 2016). Lo's (2004, 2017) Adaptive Markets Hypothesis "
    "(AMH) reconciles these facts: efficiency is not a fixed property but an evolving one, varying with the "
    "population of participants, their biases and the competition among their strategies.")
para(
    "The AMH is descriptive. It predicts that return predictability should wax and wane, a prediction supported "
    "by sub-period evidence (Kim, Shamsuddin & Lim, 2011; Urquhart & Hudson, 2013), but it does not specify how a "
    "participant could recognize, in real time and before trading, that the market has entered a less "
    "efficient state. Without such a mechanism the hypothesis cannot guide a decision. This paper asks whether "
    "a transparent, literature-grounded mechanism of this kind works out of sample.")
para(
    "We construct a three-part system, which we call Temporal Behavioral Market Dynamics (TBMD): a Real-Time "
    "Efficiency Score (RES) built from three standard random-walk diagnostics; a Behavioral Proxy Composite "
    "(BPC) of five observable proxies for the behavioral activity the AMH emphasizes; and a Regime-Conditional "
    "Filter (RCF) that trades a cross-sectional signal only when the RES indicates inefficiency. The central "
    "testable claim is that conditioning on the RES improves the risk-adjusted performance of the underlying "
    "signal relative to deploying it unconditionally.")
para(
    "Because it is easy to find a configuration of such a system that looks good in sample (Harvey, Liu & Zhu, "
    "2016; Bailey & López de Prado, 2014), the design of the test matters as much as the system. We fix a "
    "grid of twelve configurations in advance, select one using only 2005–2020 data, deflate its Sharpe "
    "ratio for the selection, and then evaluate that single configuration once on a 2021–2025 lockbox "
    "period that played no role in selecting it. We assess regime detection against four independent "
    "definitions of market stress rather than one, repeat the exercise on German equities without re-tuning, "
    "and vary transaction costs.")
para("The paper makes four contributions:")
bullet('A fully specified operationalization.',
       "Every component of the RES, BPC and RCF is defined by a closed-form rule on daily data, with all windows, "
       "thresholds and the selection rule stated in advance, and all code is released.")
bullet('Evidence on regime detection.',
       f"The RES does not identify stress regimes in U.S. or German data (AUC between {f(auc_range[0], 2)} and "
       f"{f(auc_range[1], 2)} in the U.S.; confidence intervals cover 0.5 in five of six tests) and is weak even "
       "in a simulation designed to contain the regimes it targets.")
bullet('Evidence on regime-conditional trading.',
       f"The selected strategy's lockbox Sharpe ratio is {f(rcfL['sharpe'], 2)}, and the difference between "
       f"conditional and unconditional deployment is not statistically distinguishable from zero on the lockbox "
       f"(p = {f(dL['p_two_sided'], 2)}), over the full U.S. sample (p = {f(dF['p_two_sided'], 2)}) or in "
       f"Germany (p = {f(de_diff['p_two_sided'], 2)}).")
bullet('A decomposition of what does and does not carry information.',
       "Among the behavioral proxies, aggregate illiquidity is strongly and non-circularly associated with "
       "market stress; the VIX-based sentiment proxy is strongly associated only in a partly mechanical sense; "
       "and the autocorrelation-based diagnostics on which the RES rests carry essentially none.")
para(
    "Section II reviews the literature. Section III defines the framework and Section IV the data and research "
    "design. Section V reports results, Section VI discusses them and their limitations, and Section VII "
    "concludes.")

# ------------------------------------------------------------------
# II. Literature
# ------------------------------------------------------------------
h1('II. Related Literature')
h2('A. Market Efficiency and Its Limits')
para(
    "Fama (1970) set out the weak, semi-strong and strong forms of the Efficient Market Hypothesis, which "
    "remains the null hypothesis of empirical asset pricing. Departures from it are extensively documented, "
    "including momentum (Jegadeesh & Titman, 1993), drift after earnings news (Ball & Brown, 1968) and "
    "long-horizon reversal (DeBondt & Thaler, 1985). Whether such patterns are mispricing or compensation for "
    "risk remains contested, and the replication literature urges caution: Hou et al. (2020) find that most of "
    "hundreds of published anomalies are insignificant under common methodological choices, McLean and Pontiff "
    "(2016) document post-publication decay, and Harvey et al. (2016) argue that the multiplicity of tests "
    "requires far higher significance thresholds than conventionally used.")
h2('B. The Adaptive Markets Hypothesis')
para(
    "Lo (2004, 2017) frames market efficiency in evolutionary terms: strategies earn returns while they are "
    "uncrowded and lose them as capital competes them away, so efficiency depends on the prevailing ecology of "
    "participants. The AMH predicts time-varying predictability, which has been documented using rolling and "
    "sub-period tests on long U.S. and international samples (Kim et al., 2011; Lim & Brooks, 2011; Urquhart & "
    "Hudson, 2013), and links efficiency to liquidity and trading conditions (Chordia, Roll & Subrahmanyam, "
    "2008). These studies establish that efficiency varies ex post; they do not test whether its variation can "
    "be detected ex ante and exploited, which is the question here.")
h2('C. Behavioral Proxies and Regime Models')
para(
    "Behavioral finance links market outcomes to investor biases such as loss aversion (Kahneman & Tversky, "
    "1979), herding (Christie & Huang, 1995) and sentiment (Baker & Wurgler, 2007), and shows that noise-trader "
    "risk can persist in equilibrium (De Long et al., 1990). Christie and Huang (1995) propose cross-sectional "
    "return dispersion as a herding measure and, notably, find that dispersion rises rather than falls during "
    "extreme market moves. The VIX is widely interpreted as an investor “fear gauge” (Whaley, 2000). "
    "Statistical regime models (Hamilton, 1989; Ang & Bekaert, 2002) identify latent states but are not "
    "typically tied to behavioral mechanisms or evaluated as real-time trading filters. Our design connects "
    "these strands and tests the connection out of sample.")

# ------------------------------------------------------------------
# III. Framework
# ------------------------------------------------------------------
h1('III. The TBMD Framework')
para(
    "We posit, following the AMH, that the market occupies latent regimes that differ in return "
    "autocorrelation and behavioral activity, and add one operational assertion: that the current regime can "
    "be estimated from observable data well enough to condition trading on it. The framework has three "
    "components.")
h2('A. Real-Time Efficiency Score (RES)')
para(
    "The RES measures how close the equal-weighted market return series is to a random walk using three "
    "standard diagnostics, each computed on a trailing window ending at day t: the Hurst exponent H(t) by "
    "rescaled-range analysis over 252 days (Hurst, 1951; Mandelbrot, 1971); the Lo-MacKinlay (1988) variance "
    "ratio VR(t) at lag 5 over 60 days; and the p-value of the Ljung-Box (1978) test with 10 lags over 60 days. "
    "Each is mapped to [0, 1], with 1 corresponding to a random walk:")
equation('Heff(t) = 1 − 2|H(t) − 0.5|,   VReff(t) = 1 / (1 + |VR(t) − 1|),   LBeff(t) = pLB(t)')
equation('RES(t) = [Heff(t) + VReff(t) + LBeff(t)] / 3')
para(
    "A low RES indicates departure from random-walk behavior. The inefficiency threshold τ(t) is the "
    "q-th percentile of RES over the preceding 252 trading days, so it uses no future information.")
h2('B. Behavioral Proxy Composite (BPC)')
para(
    "The BPC does not observe investor bias directly; it aggregates market-level footprints of the behavior "
    "the AMH emphasizes. Each proxy is converted to an expanding-window Z-score using only data up to day t, "
    "and the composite is their equal-weighted mean. Equal weights were chosen in advance to avoid fitting "
    "weights in sample. Table 1 defines the proxies.")
table('Table 1: Behavioral proxies',
      ['Proxy', 'Construct', 'Computation (daily, trailing)', 'Reference'],
      [['Z_herd', 'Herding', 'Negative of the cross-sectional std. dev. of returns, 20-day mean; high = low dispersion', 'Christie & Huang (1995)'],
       ['Z_VR', 'Autocorrelation', '|VR − 1| of the market return, lag 5, 60-day window', 'Lo & MacKinlay (1988)'],
       ['Z_asym', 'Loss aversion', 'Cross-sectional median of downside/upside realized volatility, 20 days', 'Ang, Chen & Xing (2006)'],
       ['Z_sent', 'Sentiment / fear', 'VIX level, 5-day mean (U.S.); DAX 21-day realized volatility (Germany)', 'Whaley (2000); Baker & Wurgler (2007)'],
       ['Z_illiq', 'Illiquidity', 'Cross-sectional median Amihud ratio |r|/dollar volume, 20-day mean', 'Amihud (2002)']],
      note='Note: each proxy is expanding-window Z-scored (minimum 252 observations) and signed so that higher '
           'values correspond to the behavioral construct named. No news-text sentiment is used.')
h2('C. Regime-Conditional Filter (RCF)')
para(
    "Because the BPC and RES are market-level series, the stocks traded are ranked by a separate per-stock "
    "composite S_i(t): the equal-weighted mean of expanding-window Z-scores of each stock's trailing m-day "
    "cumulative return (momentum; Jegadeesh & Titman, 1993), its 20-day downside/upside volatility ratio "
    "(Ang et al., 2006) and its 20-day Amihud illiquidity ratio (Amihud, 2002). Each ingredient "
    "is motivated by a documented return premium (momentum, downside risk and illiquidity, respectively). Every k trading days, using S_i on the day "
    "before rebalancing, the strategy forms an equal-weighted, dollar-neutral portfolio long the top decile "
    "and short the bottom decile of available stocks. The RCF holds this basket on day t only if")
equation('RES(t) < τ(t),')
para(
    "and is in cash otherwise. The Unconditional benchmark holds the identical basket every day. The two "
    "strategies therefore differ only in whether the RES filter is applied, which isolates the filter's "
    "contribution.")

# ------------------------------------------------------------------
# IV. Data and design
# ------------------------------------------------------------------
h1('IV. Data and Research Design')
h2('A. Data')
para(
    f"Daily split- and dividend-adjusted closing prices and share volumes for {n_us} current S&P 100 "
    f"constituents and {n_de} current DAX constituents, the SPY ETF, the S&P 500 and DAX indices and the "
    "CBOE VIX were obtained from Yahoo Finance for January 2005 to December 2025; the NBER recession "
    "indicator (USREC) was obtained from FRED. One S&P 100 ticker (BK) and one DAX ticker (1COV) had no "
    "available history and are excluded. Stocks enter the sample when their price history begins, so later "
    "listings are not excluded. One daily return that matches an unadjusted stock split (DB1.DE, 11 June "
    "2007) is set to missing by a mechanical filter; no other observation is affected. Data were downloaded "
    "on 26 September 2026.")
para(
    "Because constituents are current members, the sample is subject to survivorship bias. This biases "
    "long-only benchmarks upward; its effect on the long/short strategies is ambiguous, and it affects the RCF "
    "and Unconditional strategies equally, so it is unlikely to drive the comparison between them.")
h2('B. Measuring Market Stress')
para(
    "Real markets have no observable regime labels, so we evaluate the RES against four independent, "
    "externally defined indicators of stress, none of which enters the RES: VIX ≥ 20; 21-day realized "
    "volatility of SPY ≥ 20% annualized; an S&P 500 drawdown of at least 10% from its running peak; and "
    "an NBER recession month. For Germany we use 21-day DAX realized volatility ≥ 20% and a DAX drawdown "
    "≥ 10%. For regime-stratified reporting, days are grouped as calm (VIX < 15), transitional "
    "(15–25) and stressed (≥ 25), using DAX realized volatility with the same cut-offs for Germany.")
h2('C. Pre-Specified Selection and Lockbox Evaluation')
para(
    "The walk-forward procedure uses only information available before each trading day: Z-scores are "
    "expanding-window, the threshold τ uses the preceding 252 days, and baskets use signals dated the day "
    "before rebalancing. Transaction costs of 10 basis points are charged on each unit of one-way turnover. "
    "Before running this analysis, we fixed a grid of twelve configurations: the RES "
    "threshold percentile q ∈ {30, 40}, the rebalancing interval k ∈ {21, 63} days and the momentum "
    "window m ∈ {20, 60, 120} days. We also fixed the selection rule, choosing the configuration with the "
    "highest RCF Sharpe ratio on returns dated 2005–2020 (the development period). That configuration is "
    "evaluated once on returns dated January 2021 to December 2025 (the lockbox), and applied without change "
    "to the German sample.")
h2('D. Inference')
para(
    "We report annualized Sharpe ratios with 95% confidence intervals from a circular block bootstrap (21-day "
    "blocks; Künsch, 1989); the probabilistic Sharpe ratio (PSR) against zero, which adjusts for sample "
    "length, skewness and kurtosis (Bailey & López de Prado, 2012); and, for the selected "
    "development-period configuration, the deflated Sharpe ratio (DSR), which additionally accounts for the "
    "twelve configurations tried (Bailey & López de Prado, 2014). The difference between RCF and "
    "Unconditional Sharpe ratios is tested with a paired block bootstrap, which preserves the dependence "
    "between the two return series. Market-model alphas use Newey-West (1987) standard errors with 10 lags. "
    "Regime-detection power is summarized by the area under the ROC curve (AUC) of 1 − RES, with "
    "block-bootstrap confidence intervals, and by classification accuracy relative to the majority-class "
    "benchmark.")
h2('E. Synthetic Check')
para(
    "To check whether the RES can detect regimes in the ideal case, we apply the identical RES code to data "
    "from a three-regime Hamilton (1989) Markov-switching simulation, with 30 assets and 1,500 days. Its "
    "regimes differ in drift, volatility, market-return autocorrelation (0.02, 0.12, 0.22) and co-movement, and "
    "persist with probabilities 0.97, 0.92 and 0.88. We classify days in either of the two autocorrelated "
    "regimes as inefficient and repeat the simulation for 20 random seeds.")

# ------------------------------------------------------------------
# V. Results
# ------------------------------------------------------------------
h1('V. Results')
h2('A. Can the RES Detect Regimes It Is Designed For?')
s = synth
para(
    f"Across 20 simulated markets, the RES achieves a mean AUC of {f(s.loc['mean', 'auc'], 3)} (standard "
    f"deviation {f(s.loc['std', 'auc'], 3)}; range {f(s.loc['min', 'auc'], 2)} to {f(s.loc['max', 'auc'], 2)}). "
    f"Mean classification accuracy is {f(s.loc['mean', 'accuracy_pct'], 1)}%, below the "
    f"{f(s.loc['mean', 'majority_class_accuracy_pct'], 1)}% obtained by always predicting the majority class. "
    "Even where autocorrelation regimes exist by construction, the combination of Hurst, variance-ratio and "
    "Ljung-Box statistics over trailing windows identifies them only weakly. Regime shifts in the simulation "
    "last on average 8 to 33 days, shorter than the 60- and 252-day windows the statistics require for "
    "reliable estimation, so the diagnostics average over regimes rather than resolving them. This is an "
    "inherent tension in any window-based efficiency measure, not a matter of tuning.")
h2('B. RES Detection on Real Data')
rows = []
for name, r in us_det.iterrows():
    rows.append(['U.S.: ' + name, f(r['base_rate_pct'], 1), f(r['auc'], 3),
                 f"[{f(r['auc_ci_low'], 2)}, {f(r['auc_ci_high'], 2)}]",
                 f(r['accuracy_pct'], 1), f(r['majority_class_accuracy_pct'], 1)])
for name, r in de_det.iterrows():
    rows.append(['Germany: ' + name, f(r['base_rate_pct'], 1), f(r['auc'], 3),
                 f"[{f(r['auc_ci_low'], 2)}, {f(r['auc_ci_high'], 2)}]",
                 f(r['accuracy_pct'], 1), f(r['majority_class_accuracy_pct'], 1)])
table('Table 2: RES detection of market stress, walk-forward',
      ['Stress indicator', 'Base rate (%)', 'AUC', '95% CI', 'Accuracy (%)', 'Majority-class accuracy (%)'],
      rows,
      note=f'Note: U.S. sample {int(vix_det["n_days"]):,} days; Germany {int(de_rv["n_days"]):,} days. AUC is for '
           '1 − RES as a score for stress days; 0.5 indicates no information. Accuracy uses the walk-forward '
           'threshold at the 40th percentile. Confidence intervals are from a 21-day circular block bootstrap '
           '(500 replications).')
para(
    f"Table 2 shows the central detection result. Against all four U.S. stress indicators the RES has AUC "
    f"between {f(auc_range[0], 2)} and {f(auc_range[1], 2)}, and every confidence interval includes 0.5. Two "
    "AUCs are below 0.5, meaning the RES is, if anything, slightly higher, not lower, in stressed markets. "
    "Classification accuracy is below the majority-class benchmark in every case. In Germany, the AUC "
    f"against realized-volatility stress is {f(de_rv['auc'], 3)} with a confidence interval "
    f"({f(de_rv['auc_ci_low'], 2)} to {f(de_rv['auc_ci_high'], 2)}) that excludes 0.5, but the AUC against "
    f"drawdown stress is {f(de_dd['auc'], 3)} and accuracy is again below the majority class. We conclude "
    "that the RES carries no reliable real-time information about market stress.")
h2('C. What the Behavioral Proxies Capture')
comp_rows = []
names = {'Z_herd': 'Z_herd (herding)', 'Z_vr': 'Z_VR (autocorrelation)', 'Z_asym': 'Z_asym (loss aversion)',
         'Z_sent': 'Z_sent (VIX sentiment)', 'Z_illiq': 'Z_illiq (illiquidity)', 'BPC': 'BPC composite'}
for k, r in comp.iterrows():
    comp_rows.append([names[k]] + [f(v, 3) for v in r.iloc[:4]] + [f(r['mean | VIX >= 20'], 2, True), f(r['mean | VIX < 20'], 2, True)])
table('Table 3: Association of BPC components with market stress (U.S., full sample)',
      ['Component', 'AUC: VIX ≥ 20', 'AUC: RV ≥ 20%', 'AUC: DD ≥ 10%', 'AUC: NBER',
       'Mean | VIX ≥ 20', 'Mean | VIX < 20'],
      comp_rows,
      note='Note: AUC of each Z-scored component as a score for stress days; values below 0.5 indicate that '
           'the component is lower on stress days. RV = SPY 21-day realized volatility; DD = S&P 500 drawdown.')
para(
    f"Table 3 decomposes the BPC. Illiquidity is the most informative non-circular proxy: its AUC is "
    f"{f(zi_lo, 2)} to {f(zi_hi, 2)} across the four stress definitions, consistent with Chordia et al. "
    f"(2008) and with liquidity deteriorating in stressed markets. The VIX-based sentiment proxy has AUC "
    f"{f(zs.iloc[0], 2)} against VIX-defined stress, but this is largely mechanical, since both are functions "
    f"of the VIX. Its AUC of {f(zs.iloc[3], 2)} against NBER recessions is more informative, but it reflects "
    "that the VIX is elevated in recessions rather than a distinct behavioral signal. The herding proxy is "
    f"strongly associated with stress in the opposite direction to the herding interpretation (AUC "
    f"{f(zh.iloc[0], 2)} for VIX): cross-sectional dispersion rises in stressed markets, matching Christie and "
    "Huang's (1995) finding that dispersion widens during extreme moves. The autocorrelation and "
    "loss-aversion proxies carry little information. Because the BPC equal-weights strong and uninformative "
    "components, and includes herding with an inverted sign, the composite (AUC "
    f"{f(zb_lo, 2)}–{f(zb_hi, 2)}) is weaker than its best component. The weakness of Z_VR is "
    "consistent with the RES result, since both rest on autocorrelation diagnostics.")
h2('D. Development Period and Configuration Selection')
g_rows = [[int(r.tau_pct), int(r.refit_days), int(r.mom_window), f(r.dev_rcf_sharpe, 3), f(r.dev_unc_sharpe, 3)]
          for r in grid.itertuples()]
table('Table 4: Development-period (2005–2020) Sharpe ratios of all twelve pre-specified configurations',
      ['RES threshold q', 'Rebalance k (days)', 'Momentum m (days)', 'RCF Sharpe', 'Unconditional Sharpe'],
      g_rows,
      note='Note: annualized Sharpe ratios net of 10 bp costs. The Unconditional strategy does not depend on q.')
para(
    f"Only {n_pos} of the twelve configurations have a positive development-period RCF Sharpe ratio (Table 4). "
    f"The selection rule picks q = {sel['tau_pct']}, k = {sel['refit_days']}, m = {sel['mom_window']}, with a "
    f"Sharpe ratio of {f(dev_best, 3)}. Deflating for the twelve trials gives a DSR of {f(rcfD['dsr_n12'], 2)}, "
    "far below conventional thresholds for a genuine effect. Its PSR against zero is "
    f"{f(rcfD['psr_vs_0'], 2)} and the mean-return t-statistic is {f(rcfD['t_stat_mean'], 2)}. Even before the "
    "lockbox is opened, the development-period evidence for the strategy is weak.")
h2('E. Lockbox Performance')


def perf_row(label, r, alpha=True):
    return [label, f(r['ann_return_pct'], 1, True), f(r['ann_vol_pct'], 1), f(r['sharpe'], 2),
            f"[{f(r['sharpe_ci_low'], 2)}, {f(r['sharpe_ci_high'], 2)}]", f(r['psr_vs_0'], 2),
            f(r['max_drawdown_pct'], 1), f(r['pct_days_active'], 0),
            (f"{f(r['alpha_ann_pct'], 1, True)} ({f(r['alpha_t_nw'], 2)})" if alpha and pd.notna(r.get('alpha_ann_pct')) else '—')]


hdr = ['Strategy', 'Ann. return (%)', 'Ann. vol (%)', 'Sharpe', '95% CI', 'PSR', 'Max DD (%)',
       'Days active (%)', 'Alpha % (t)']
table('Table 5: Lockbox performance, 2021–2025 (selected configuration, U.S.)',
      hdr,
      [perf_row('RCF', rcfL), perf_row('Unconditional', uncL), perf_row('Buy & hold, equal weight', ewL),
       perf_row('Buy & hold, SPY', spyL, alpha=False)],
      note='Note: net of 10 bp one-way costs. PSR = probabilistic Sharpe ratio against zero. Alpha is the '
           'annualized intercept of a regression on SPY returns, with Newey-West (10-lag) t-statistics.')
para(
    f"On the lockbox the selected RCF loses {f(-rcfL['ann_return_pct'], 1)}% per year with a Sharpe ratio of "
    f"{f(rcfL['sharpe'], 2)}. Its confidence interval ({f(rcfL['sharpe_ci_low'], 2)} to "
    f"{f(rcfL['sharpe_ci_high'], 2)}) lies entirely below zero, and its market-model alpha is "
    f"{f(rcfL['alpha_ann_pct'], 1)}% (t = {f(rcfL['alpha_t_nw'], 2)}). The Unconditional strategy also loses "
    f"money (Sharpe {f(uncL['sharpe'], 2)}), while passive benchmarks earn Sharpe ratios near one. Performance "
    "therefore deteriorates from development to lockbox, as expected when a configuration selected on noise "
    "meets new data.")
h2('F. Does the Filter Add Value?')
table('Table 6: Sharpe ratio of RCF minus Unconditional (selected configuration)',
      ['Sample', 'Difference', '95% CI', 'p-value (two-sided)'],
      [['U.S. development, 2005–2020', f(dD['sharpe_diff'], 3, True), f"[{f(dD['ci_low'], 2)}, {f(dD['ci_high'], 2)}]", f(dD['p_two_sided'], 3)],
       ['U.S. lockbox, 2021–2025', f(dL['sharpe_diff'], 3, True), f"[{f(dL['ci_low'], 2)}, {f(dL['ci_high'], 2)}]", f(dL['p_two_sided'], 3)],
       ['U.S. full sample', f(dF['sharpe_diff'], 3, True), f"[{f(dF['ci_low'], 2)}, {f(dF['ci_high'], 2)}]", f(dF['p_two_sided'], 3)],
       ['Germany, full sample (no re-tuning)', f(de_diff['sharpe_diff'], 3, True), f"[{f(de_diff['ci_low'], 2)}, {f(de_diff['ci_high'], 2)}]", f(de_diff['p_two_sided'], 3)]],
      note='Note: paired circular block bootstrap (21-day blocks, 2,000 replications) of the difference in '
           'annualized Sharpe ratios. The development-period figure is for the configuration selected on that '
           'period and is therefore biased upward.')
para(
    "The paper's core hypothesis is that the RCF improves on the Unconditional strategy. Table 6 does not "
    f"support it. The development-period difference ({f(dD['sharpe_diff'], 2, True)}) is inflated by "
    f"selection and still not significant at 5%. It reverses sign on the lockbox ({f(dL['sharpe_diff'], 2, True)}) "
    f"and in Germany ({f(de_diff['sharpe_diff'], 2, True)}), and none of the four differences is significant. "
    "This follows directly from Section V.B: a filter that cannot distinguish stressed from calm markets "
    "cannot systematically improve a strategy by switching it on and off.")
h2('G. Performance by Regime')
rb_rows = [[r.regime, int(r.n_days), f(r.pct_time, 1), f(r.rcf_ann_return_pct, 1, True), f(r.rcf_sharpe, 2),
            f(r.unconditional_ann_return_pct, 1, True), f(r.unconditional_sharpe, 2), f(r.buyhold_ew_ann_return_pct, 1, True)]
           for r in rb.itertuples()]
table('Table 7: Annualized returns by VIX regime, U.S. full sample (selected configuration)',
      ['VIX regime', 'Days', '% time', 'RCF return (%)', 'RCF Sharpe', 'Uncond. return (%)', 'Uncond. Sharpe',
       'Buy & hold return (%)'],
      rb_rows,
      note='Note: regimes are for reporting only and are not used by any strategy.')
para(
    f"Both long/short strategies are close to market-neutral (full-sample betas of {f(rcfF['beta'], 2)} and "
    f"{f(uncF['beta'], 2)}), so comparing them with buy-and-hold by regime mainly reflects market exposure: "
    f"they avoid the {f(stress.buyhold_ew_ann_return_pct, 1)}% annualized loss of the equal-weighted market in "
    "stressed periods and forgo its large gains in calm ones. The relevant comparison is between RCF and "
    f"Unconditional. In stressed U.S. periods the RCF loses {f(-stress.rcf_ann_return_pct, 1)}% against "
    f"{f(-stress.unconditional_ann_return_pct, 1)}% for the Unconditional strategy, mostly because it trades "
    "less. The same comparison reverses in Germany, where the RCF earns "
    f"{f(de_stress.rcf_ann_return_pct, 1, True)}% against {f(de_stress.unconditional_ann_return_pct, 1, True)}% "
    "in high-volatility periods. We find no consistent regime in which the filter adds value.")
h2('H. Robustness')
c_rows = []
for bps in sorted(cost.cost_bps.unique()):
    r = cost[(cost.cost_bps == bps) & (cost.strategy == 'RCF')].iloc[0]
    u = cost[(cost.cost_bps == bps) & (cost.strategy == 'Unconditional')].iloc[0]
    c_rows.append([f'{bps} bp', f(r.full_sharpe, 3), f(r.lockbox_sharpe, 3), f(u.full_sharpe, 3), f(u.lockbox_sharpe, 3)])
table('Table 8: Sensitivity of Sharpe ratios to transaction costs (U.S., selected configuration)',
      ['One-way cost', 'RCF full', 'RCF lockbox', 'Uncond. full', 'Uncond. lockbox'], c_rows)
para(
    f"Transaction costs do not explain the results. With zero costs the RCF's lockbox Sharpe ratio is still "
    f"{f(cost0.lockbox_sharpe, 2)} (Table 8), so the losses come from the signal rather than from trading "
    "frictions.")
table(f'Table 9: Out-of-U.S. replication, DAX constituents, full sample (no re-tuning)',
      hdr,
      [perf_row('RCF', de_rcf), perf_row('Unconditional', de_unc), perf_row('Buy & hold, equal weight', de_ew)],
      note='Note: selected U.S. configuration applied unchanged. Alpha is relative to the DAX index.')
para(
    f"Applied unchanged to {n_de} DAX constituents (Table 9), the RCF has a Sharpe ratio of {f(de_rcf['sharpe'], 2)} "
    f"and the Unconditional strategy {f(de_unc['sharpe'], 2)}. Neither is significantly different from zero, and "
    "the filter again fails to improve on unconditional deployment. The conclusions are not specific to the "
    "U.S. market.")

# ------------------------------------------------------------------
# VI. Discussion
# ------------------------------------------------------------------
h1('VI. Discussion')
h2('A. Interpretation')
para(
    "The evidence points to a specific failure. The system's premise is that departures from random-walk "
    "behavior in the market index identify behaviorally driven regimes. In the data, the three diagnostics "
    "that define the RES do not move with any of four independent measures of market stress, and they "
    "identify even engineered autocorrelation regimes only weakly because those regimes are short relative to "
    "the estimation windows. Since the trading filter inherits the detector's information, its failure "
    "follows. By contrast, a market-level quantity the AMH links to efficiency, liquidity, does move strongly "
    "with stress, and it does so without being constructed from the stress measures themselves.")
para(
    "Nothing here refutes the AMH, and we do not claim that market efficiency is constant. The results "
    "concern one natural class of operationalizations: window-based, index-level random-walk diagnostics used "
    "as a real-time on/off switch. They suggest that detecting the AMH's regimes in real time requires "
    "measures closer to the mechanisms the hypothesis emphasizes, such as liquidity, funding and the "
    "competitive intensity of arbitrage capital, rather than measures of statistical predictability, which "
    "are noisy at the horizons over which regimes appear to change.")
h2('B. Limitations')
bullet('Survivorship.',
       "Constituents are current members, and point-in-time membership with delisted-stock returns (e.g., CRSP) "
       "was not available. This biases the passive benchmarks upward but affects the RCF and Unconditional "
       "strategies equally, so it is unlikely to drive the paper's central comparison.")
bullet('Scope of the configuration space.',
       "Twelve configurations were examined, all sharing the same diagnostics, proxies and portfolio "
       "construction. Other operationalizations, such as fitted proxy weights, different diagnostics or "
       "intraday data, are not ruled out. By design, we did not search for one after seeing the lockbox.")
bullet('Sentiment measurement.',
       "The sentiment proxy is VIX-based in the U.S. and volatility-based in Germany. It is not a text-based or "
       "survey measure, and it is partly mechanical with respect to volatility-defined stress, as documented in "
       "Table 3.")
bullet('Lockbox purity.',
       "The configuration was selected using development-period data only. However, exploratory runs made "
       "while building the pipeline (a single configuration on a 27-stock subset, 2005–2023) overlapped "
       "2021–2023. They did not inform the grid or the selection rule, but the lockbox is not pristine in "
       "the strictest sense, and the German sample, to which no design choice was fitted, provides the "
       "cleaner out-of-sample check.")
bullet('Sample.',
       "The lockbox covers five years. Its conclusions agree with the full-sample and German evidence but rest "
       "on a single post-2020 period.")
h2('C. Future Research')
para(
    "Three extensions follow directly. First, liquidity-based detectors, which the component analysis "
    "identifies as informative, could be evaluated as regime filters under the same pre-specified protocol. "
    "Second, the tension between regime persistence and estimation-window length could be addressed with "
    "higher-frequency data, which would allow the same diagnostics to be estimated over shorter calendar "
    "windows. Third, the analysis should be repeated on point-in-time CRSP constituents.")

# ------------------------------------------------------------------
# VII. Conclusion
# ------------------------------------------------------------------
h1('VII. Conclusion')
para(
    "We asked whether the Adaptive Markets Hypothesis can be turned into a real-time decision rule, building "
    "a transparent system and testing it under a pre-specified selection rule, a held-out lockbox, four "
    "independent regime definitions, a second national market and a range of transaction costs. In this "
    "specification it cannot. The efficiency score does not detect market stress and "
    f"barely detects regimes in a simulation built to contain them. The selected strategy's Sharpe ratio falls "
    f"from {f(rcfD['sharpe'], 2)} in development to {f(rcfL['sharpe'], 2)} on the lockbox, and conditioning on "
    "the score does not reliably improve on deploying the same signal unconditionally in either market. What "
    "survives is narrower: aggregate illiquidity is a strong, non-circular correlate of market stress, and "
    "the autocorrelation diagnostics commonly used to measure time-varying efficiency are not. We hope the "
    "framework, code and data released with this paper will be useful for testing better candidates under "
    "the same discipline.")

h1('Data and Code Availability')
para(
    "All data are publicly available (Yahoo Finance; FRED) and were downloaded on 26 September 2026. Because "
    "Yahoo Finance data cannot be redistributed, they are not included; re-downloading may yield slightly "
    "different adjusted prices. The download script (build_dataset.py), the full "
    "analysis (run_paper_analysis.py) and the script that generates every table and in-text number in this "
    "manuscript from the analysis outputs (build_paper.py) are available in the author's repository.")

h1('References')
refs = [
    "Amihud, Y. (2002). Illiquidity and stock returns: Cross-section and time-series effects. Journal of Financial Markets, 5(1), 31–56.",
    "Ang, A., & Bekaert, G. (2002). International asset allocation with regime shifts. Review of Financial Studies, 15(4), 1137–1187.",
    "Ang, A., Chen, J., & Xing, Y. (2006). Downside risk. Review of Financial Studies, 19(4), 1191–1239.",
    "Bailey, D. H., & López de Prado, M. (2012). The Sharpe ratio efficient frontier. Journal of Risk, 15(2), 3–44.",
    "Bailey, D. H., & López de Prado, M. (2014). The deflated Sharpe ratio: Correcting for selection bias, backtest overfitting, and non-normality. Journal of Portfolio Management, 40(5), 94–107.",
    "Baker, M., & Wurgler, J. (2007). Investor sentiment in the stock market. Journal of Economic Perspectives, 21(2), 129–151.",
    "Ball, R., & Brown, P. (1968). An empirical evaluation of accounting income numbers. Journal of Accounting Research, 6(2), 159–178.",
    "Chordia, T., Roll, R., & Subrahmanyam, A. (2008). Liquidity and market efficiency. Journal of Financial Economics, 87(2), 249–268.",
    "Christie, W. G., & Huang, R. D. (1995). Following the pied piper: Do individual returns herd around the market? Financial Analysts Journal, 51(4), 31–37.",
    "DeBondt, W. F. M., & Thaler, R. (1985). Does the stock market overreact? Journal of Finance, 40(3), 793–805.",
    "De Long, J. B., Shleifer, A., Summers, L. H., & Waldmann, R. J. (1990). Noise trader risk in financial markets. Journal of Political Economy, 98(4), 703–738.",
    "Fama, E. F. (1970). Efficient capital markets: A review of theory and empirical work. Journal of Finance, 25(2), 383–417.",
    "Hamilton, J. D. (1989). A new approach to the economic analysis of nonstationary time series and the business cycle. Econometrica, 57(2), 357–384.",
    "Harvey, C. R., Liu, Y., & Zhu, H. (2016). …and the cross-section of expected returns. Review of Financial Studies, 29(1), 5–68.",
    "Hou, K., Xue, C., & Zhang, L. (2020). Replicating anomalies. Review of Financial Studies, 33(5), 2019–2133.",
    "Hurst, H. E. (1951). Long-term storage capacity of reservoirs. Transactions of the American Society of Civil Engineers, 116, 770–799.",
    "Jegadeesh, N., & Titman, S. (1993). Returns to buying winners and selling losers: Implications for stock market efficiency. Journal of Finance, 48(1), 65–91.",
    "Kahneman, D., & Tversky, A. (1979). Prospect theory: An analysis of decision under risk. Econometrica, 47(2), 263–291.",
    "Kim, J. H., Shamsuddin, A., & Lim, K.-P. (2011). Stock return predictability and the adaptive markets hypothesis: Evidence from century-long U.S. data. Journal of Empirical Finance, 18(5), 868–879.",
    "Künsch, H. R. (1989). The jackknife and the bootstrap for general stationary observations. Annals of Statistics, 17(3), 1217–1241.",
    "Lim, K.-P., & Brooks, R. (2011). The evolution of stock market efficiency over time: A survey of the empirical literature. Journal of Economic Surveys, 25(1), 69–108.",
    "Ljung, G. M., & Box, G. E. P. (1978). On a measure of lack of fit in time series models. Biometrika, 65(2), 297–303.",
    "Lo, A. W. (2004). The adaptive markets hypothesis: Market efficiency from an evolutionary perspective. Journal of Portfolio Management, 30(5), 15–29.",
    "Lo, A. W. (2017). Adaptive Markets: Financial Evolution at the Speed of Thought. Princeton University Press.",
    "Lo, A. W., & MacKinlay, A. C. (1988). Stock market prices do not follow random walks: Evidence from a simple specification test. Review of Financial Studies, 1(1), 41–66.",
    "Mandelbrot, B. B. (1971). When can price be arbitraged efficiently? A limit to the validity of the random walk and martingale models. Review of Economics and Statistics, 53(3), 225–236.",
    "McLean, R. D., & Pontiff, J. (2016). Does academic research destroy stock return predictability? Journal of Finance, 71(1), 5–32.",
    "Newey, W. K., & West, K. D. (1987). A simple, positive semi-definite, heteroskedasticity and autocorrelation consistent covariance matrix. Econometrica, 55(3), 703–708.",
    "Urquhart, A., & Hudson, R. (2013). Efficient or adaptive markets? Evidence from major stock markets using very long run historic data. International Review of Financial Analysis, 28, 130–142.",
    "Whaley, R. E. (2000). The investor fear gauge. Journal of Portfolio Management, 26(3), 12–17.",
]
for r in refs:
    p = para(r, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)

doc.save(DOCX)
print('Wrote', DOCX)
